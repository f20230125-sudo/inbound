"""Synthetic document bytes, independent of extraction; fixed metadata and ZIP timestamps."""

import io
import zipfile
from datetime import UTC, datetime

from docx import Document
from reportlab.pdfgen import canvas


def pdf_bytes(text, *, multipage=False, image_only=False):
    stream = io.BytesIO()
    pdf = canvas.Canvas(stream, invariant=1, pageCompression=0)
    pdf.setTitle("Fictional internship description")
    pdf.setAuthor("Synthetic fixture generator")
    pdf.setCreator("Inbound synthetic fixtures")
    if image_only:
        # A raster image with no embedded text layer; no OCR is supported.
        from PIL import Image
        from reportlab.lib.utils import ImageReader

        image = Image.new("RGB", (200, 50), "white")
        pdf.drawImage(ImageReader(image), 40, 700)
    else:
        lines = text.splitlines()
        if multipage:
            # Split a requirement across the page boundary without changing its authored text.
            text = text.replace("SQLite FTS5", "SQLite\nFTS5")
            lines = text.splitlines()
        y = 780
        for line in lines:
            if multipage and line == "FTS5; BM25":
                pdf.showPage()
                y = 780
            pdf.drawString(40, y, line)
            y -= 18
    pdf.save()
    return stream.getvalue()


def docx_bytes(text, layout):
    document = Document()
    document.core_properties.created = datetime(2026, 10, 28, tzinfo=UTC)
    document.core_properties.modified = datetime(2026, 10, 28, tzinfo=UTC)
    document.core_properties.author = "Synthetic fixture generator"
    document.core_properties.title = "Fictional internship description"
    lines = text.splitlines()
    if "merged header" in layout:
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).merge(table.cell(0, 1)).text = "Fictional role requirements"
        for i in range(0, len(lines), 2):
            cells = table.add_row().cells
            # Restarted numbering in every row; role identity comes from the title.
            cells[0].text = lines[i].replace(f"({i // 2 + 1})", "(1)")
            cells[1].text = lines[i + 1]
    elif "paragraphs plus table" in layout:
        for line in lines[:2]:
            document.add_paragraph(line)
        table = document.add_table(rows=1, cols=1)
        table.cell(0, 0).text = "\n".join(lines[2:])
    elif "DOCX table" in layout:
        table = document.add_table(rows=len(lines), cols=1)
        for row, line in zip(table.rows, lines, strict=True):
            row.cells[0].text = line
    else:
        for line in lines:
            document.add_paragraph(line)
    stream = io.BytesIO()
    document.save(stream)
    output = io.BytesIO()
    with (
        zipfile.ZipFile(io.BytesIO(stream.getvalue())) as source,
        zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as target,
    ):
        for name in sorted(source.namelist()):
            info = zipfile.ZipInfo(name, date_time=(2026, 10, 28, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            target.writestr(info, source.read(name))
    return output.getvalue()


def attachment_bytes(case, attachment):
    status = next(
        a["parse_status_m2"] for a in case["gold"]["attachments"] if a["name"] == attachment["name"]
    )
    if status == "missing":
        return None
    layout = case["attachment_plan"].get("layout", "")
    if status == "unreadable" and "Image-only" not in layout:
        return b"INBOUND synthetic corrupt document\n"
    text = case["attachment_plan"]["text"]
    if "brochure" in attachment["name"]:
        text = "Fictional company brochure. No role requirements."
    if attachment["mime"] == "application/pdf":
        return pdf_bytes(text, multipage="Two-page" in layout, image_only="Image-only" in layout)
    return docx_bytes(text, layout)
